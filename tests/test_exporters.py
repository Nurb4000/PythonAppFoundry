"""Tests for multi-format export functionality."""
import io
import json
import csv
import pytest


class FakeRow:
    """Minimal stand-in for a SQLAlchemy model row."""
    def __init__(self, **kwargs):
        self.__dict__.update(kwargs)


class TestSerializeValue:
    """Test _serialize_value helper."""

    def test_none(self):
        from app.services.exporters import _serialize_value
        assert _serialize_value(None) is None

    def test_int(self):
        from app.services.exporters import _serialize_value
        assert _serialize_value(42) == 42

    def test_string(self):
        from app.services.exporters import _serialize_value
        assert _serialize_value('hello') == 'hello'

    def test_datetime(self):
        from datetime import datetime
        from app.services.exporters import _serialize_value
        dt = datetime(2024, 1, 15, 10, 30, 0)
        result = _serialize_value(dt)
        assert isinstance(result, str)
        assert '2024-01-15' in result

    def test_date(self):
        from datetime import date
        from app.services.exporters import _serialize_value
        d = date(2024, 1, 15)
        result = _serialize_value(d)
        assert result == '2024-01-15'


class TestExportJson:
    """Test JSON export."""

    def test_basic_export(self):
        from app.services.exporters import _export_json
        rows = [FakeRow(id=1, name='alice'), FakeRow(id=2, name='bob')]
        resp = _export_json('users', ['id', 'name'], rows, False)

        assert resp.mimetype == 'application/json'
        assert 'attachment' in resp.headers.get('Content-Disposition', '')

        data = json.loads(resp.data.decode())
        assert data['name'] == 'users'
        assert data['columns'] == ['id', 'name']
        assert len(data['rows']) == 2
        assert data['rows'][0]['name'] == 'alice'


class TestExportCsv:
    """Test CSV export."""

    def test_basic_export(self):
        from app.services.admin_utils import _export_csv
        rows = [FakeRow(id=1, name='alice')]
        resp = _export_csv('users', ['id', 'name'], rows, False)

        assert resp.mimetype == 'text/csv'
        reader = csv.reader(io.StringIO(resp.data.decode()))
        headers = next(reader)
        assert headers == ['id', 'name']
        data_row = next(reader)
        assert data_row == ['1', 'alice']


class TestExportXlsx:
    """Test XLSX export."""

    def test_basic_export(self):
        from app.services.exporters import _export_xlsx, HAS_OPENPYXL
        if not HAS_OPENPYXL:
            pytest.skip('openpyxl not installed')

        rows = [FakeRow(id=1, name='alice'), FakeRow(id=2, name='bob')]
        resp = _export_xlsx('users', ['id', 'name'], rows, False)

        assert resp.status_code == 200
        assert 'application/vnd.openxmlformats' in resp.mimetype

        buf = io.BytesIO(resp.data)
        import openpyxl
        wb = openpyxl.load_workbook(buf)
        ws = wb.active
        assert ws.cell(1, 1).value == 'id'
        assert ws.cell(1, 2).value == 'name'
        assert ws.cell(2, 1).value == 1
        assert ws.cell(2, 2).value == 'alice'


class TestExportPdf:
    """Test PDF export."""

    def test_basic_export(self):
        from app.services.exporters import _export_pdf, HAS_REPORTLAB
        if not HAS_REPORTLAB:
            pytest.skip('reportlab not installed')

        rows = [FakeRow(id=1, name='alice'), FakeRow(id=2, name='bob')]
        resp = _export_pdf('users', ['id', 'name'], rows, False)

        assert resp.status_code == 200
        assert resp.mimetype == 'application/pdf'
        assert resp.data[:4] == b'%PDF'


class TestExportFormatsRegistry:
    """Test the EXPORT_FORMATS registry."""

    def test_registered_formats(self):
        from app.services.exporters import EXPORT_FORMATS
        assert 'json' in EXPORT_FORMATS
        assert 'xlsx' in EXPORT_FORMATS
        assert 'pdf' in EXPORT_FORMATS
        assert len(EXPORT_FORMATS) == 3


class TestRowToDict:
    """Test _row_to_dict helper."""

    def test_with_module(self):
        from app.services.exporters import _row_to_dict
        mod = FakeRow(name='mymod')
        row = FakeRow(id=1, name='test', module=mod)
        d = _row_to_dict(row, ['id', 'name'], has_module=True)
        assert d['module'] == 'mymod'
        assert d['id'] == 1
        assert d['name'] == 'test'

    def test_without_module(self):
        from app.services.exporters import _row_to_dict
        row = FakeRow(id=1, name='test')
        d = _row_to_dict(row, ['id', 'name'], has_module=False)
        assert 'module' not in d
        assert d['id'] == 1
